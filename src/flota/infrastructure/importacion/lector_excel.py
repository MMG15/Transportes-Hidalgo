from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

from openpyxl import load_workbook

from flota.application.dto import ErrorLectura, FilaCarga
from flota.application.ports import LectorCargas
from flota.infrastructure.importacion.mapeo import procesar_filas


class LectorExcel(LectorCargas):
    """Lee la primera hoja de un .xlsx. La fila 1 debe tener los encabezados."""

    def leer(self, ruta: Path) -> Iterator[FilaCarga | ErrorLectura]:
        # read_only: no carga todo el archivo en memoria (sirve para planillas grandes).
        # data_only: devuelve el resultado de las fórmulas, no la fórmula.
        libro = load_workbook(ruta, read_only=True, data_only=True)
        try:
            filas = libro.worksheets[0].iter_rows(values_only=True)
            encabezado = next(filas, None)
            if encabezado is None:
                yield ErrorLectura(1, "El archivo está vacío")
                return
            yield from procesar_filas(encabezado, enumerate(filas, start=2))
        finally:
            libro.close()
