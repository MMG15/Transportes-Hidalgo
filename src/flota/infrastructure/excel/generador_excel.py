"""Reporte de consumo en Excel con la identidad de Transportes Hidalgo:
portada con logo, encabezados en azul de marca, filtros y alertas resaltadas en naranja."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import BinaryIO

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from flota.application.dto import ReporteConsumo
from flota.application.ports import GeneradorReporte
from flota.marca import NOMBRE_EMPRESA, color_excel

_ENCABEZADO_FONDO = PatternFill("solid", fgColor=color_excel("azul"))
_ENCABEZADO_FUENTE = Font(bold=True, color="FFFFFF")
_ENCABEZADO_BORDE = Border(bottom=Side(style="medium", color=color_excel("naranja")))
_ALERTA_FONDO = PatternFill("solid", fgColor=color_excel("naranja-claro"))
_ALERTA_FUENTE = Font(bold=True, color=color_excel("naranja-oscuro"))
_ACENTO_FONDO = PatternFill("solid", fgColor=color_excel("naranja"))
_FORMATO_FECHA = "dd/mm/yyyy hh:mm"
_FORMATO_NUM = "#,##0.00"
_FORMATO_PCT = '0.0"%"'
_ALTO_LOGO_PX = 60

COLUMNAS_PLANTILLA = [
    "fecha",
    "patente",
    "litros",
    "precio_litro",
    "odometro_km",
    "estacion",
    "referencia",
]

type Columnas = list[tuple[str, int, str | None]]  # (título, ancho, formato numérico)


class GeneradorReporteExcel(GeneradorReporte):
    def __init__(self, logo: Path | None = None) -> None:
        self._logo = logo

    def generar(self, reporte: ReporteConsumo, destino: Path | BinaryIO) -> None:
        libro = Workbook()
        self._hoja_resumen(libro.active, reporte)
        self._hoja_viajes(libro.create_sheet("Consumo por viaje"), reporte)
        libro.save(destino)

    def _hoja_resumen(self, hoja: Worksheet, reporte: ReporteConsumo) -> None:
        hoja.title = "Resumen por camión"
        hoja.sheet_view.showGridLines = False
        columnas: Columnas = [
            ("Patente", 12, None),
            ("Viajes", 9, None),
            ("Km", 12, _FORMATO_NUM),
            ("Litros", 12, _FORMATO_NUM),
            ("Costo $", 15, _FORMATO_NUM),
            ("L/100km promedio", 16, _FORMATO_NUM),
            ("Ref. L/100km", 13, _FORMATO_NUM),
            ("Desvío", 10, _FORMATO_PCT),
            ("Viajes con alerta", 16, None),
        ]
        fila_tabla = self._portada(hoja, reporte, len(columnas))
        _escribir_encabezado(hoja, columnas, fila_tabla)
        for r in reporte.resumen_por_camion:
            hoja.append(
                [
                    r.patente,
                    r.viajes,
                    r.km_recorridos,
                    r.litros,
                    r.costo,
                    r.consumo_promedio_l100km,
                    r.consumo_referencia_l100km,
                    r.desvio_pct,
                    r.viajes_con_alerta,
                ]
            )
            if r.viajes_con_alerta:
                hoja.cell(hoja.max_row, len(columnas)).font = _ALERTA_FUENTE
        _aplicar_formatos(hoja, columnas, fila_tabla)

    def _portada(self, hoja: Worksheet, reporte: ReporteConsumo, ancho: int) -> int:
        """Logo, título y período. Devuelve la fila donde empieza la tabla."""
        if self._logo and self._logo.exists():
            imagen = Image(str(self._logo))
            imagen.width = round(imagen.width * _ALTO_LOGO_PX / imagen.height)
            imagen.height = _ALTO_LOGO_PX
            hoja.add_image(imagen, "A1")
        hoja["A5"] = "Reporte de consumo de combustible"
        hoja["A5"].font = Font(bold=True, size=16, color=color_excel("azul-oscuro"))
        hoja["A6"] = (
            f"{NOMBRE_EMPRESA} · Período {reporte.desde:%d/%m/%Y} al {reporte.hasta:%d/%m/%Y}"
            f" · Generado el {datetime.now():%d/%m/%Y %H:%M}"
        )
        hoja["A6"].font = Font(color=color_excel("texto-suave"))
        for columna in range(1, ancho + 1):  # franja naranja de acento
            hoja.cell(7, columna).fill = _ACENTO_FONDO
        hoja.row_dimensions[7].height = 4
        return 9

    def _hoja_viajes(self, hoja: Worksheet, reporte: ReporteConsumo) -> None:
        columnas: Columnas = [
            ("Salida", 17, _FORMATO_FECHA),
            ("Llegada", 17, _FORMATO_FECHA),
            ("Estado", 11, None),
            ("Patente", 10, None),
            ("Chofer", 22, None),
            ("Cliente", 28, None),
            ("Producto", 16, None),
            ("Cantidad", 10, _FORMATO_NUM),
            ("Unidad", 7, None),
            ("Origen", 18, None),
            ("Destino", 18, None),
            ("Km", 10, _FORMATO_NUM),
            ("Litros", 10, _FORMATO_NUM),
            ("Costo $", 13, _FORMATO_NUM),
            ("L/100km", 10, _FORMATO_NUM),
            ("Ref. L/100km", 12, _FORMATO_NUM),
            ("Desvío", 9, _FORMATO_PCT),
            ("L por unidad", 12, _FORMATO_NUM),
            ("Alerta", 8, None),
        ]
        _escribir_encabezado(hoja, columnas, 1)
        for f in reporte.filas:
            hoja.append(
                [
                    f.fecha_salida,
                    f.fecha_llegada,
                    f.estado,
                    f.patente,
                    f.chofer,
                    f.cliente,
                    f.producto,
                    f.cantidad,
                    f.unidad,
                    f.origen,
                    f.destino,
                    f.km_recorridos,
                    f.litros,
                    f.costo,
                    f.consumo_l100km,
                    f.consumo_referencia_l100km,
                    f.desvio_pct,
                    f.litros_por_unidad,
                    "SÍ" if f.alerta else "",
                ]
            )
            if f.alerta:
                for celda in hoja[hoja.max_row]:
                    celda.fill = _ALERTA_FONDO
                hoja.cell(hoja.max_row, len(columnas)).font = _ALERTA_FUENTE
        _aplicar_formatos(hoja, columnas, 1)


def _escribir_encabezado(hoja: Worksheet, columnas: Columnas, fila: int) -> None:
    for indice, (titulo, _, _) in enumerate(columnas, start=1):
        celda = hoja.cell(fila, indice, titulo)
        celda.fill = _ENCABEZADO_FONDO
        celda.font = _ENCABEZADO_FUENTE
        celda.border = _ENCABEZADO_BORDE
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hoja.row_dimensions[fila].height = 30
    hoja.freeze_panes = f"A{fila + 1}"  # el encabezado queda fijo al hacer scroll


def _aplicar_formatos(hoja: Worksheet, columnas: Columnas, fila_encabezado: int) -> None:
    for indice, (_, ancho, formato) in enumerate(columnas, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
        if formato:
            for (celda,) in hoja.iter_rows(
                min_row=fila_encabezado + 1, min_col=indice, max_col=indice
            ):
                celda.number_format = formato
    ultima = get_column_letter(len(columnas))
    hoja.auto_filter.ref = f"A{fila_encabezado}:{ultima}{max(hoja.max_row, fila_encabezado)}"


def crear_plantilla_importacion(destino: Path | BinaryIO) -> None:
    """Planilla vacía con las columnas que entiende el importador, más una fila de ejemplo."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Cargas"
    columnas: Columnas = [(c, 16, None) for c in COLUMNAS_PLANTILLA]
    _escribir_encabezado(hoja, columnas, 1)
    hoja.append(["21/10/2026 08:30", "AB123CD", 350.5, 1450.0, 125400, "YPF Ruta 5", "T-0001"])
    _aplicar_formatos(hoja, columnas, 1)
    libro.save(destino)
